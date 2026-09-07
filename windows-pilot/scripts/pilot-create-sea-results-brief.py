from __future__ import annotations

from collections import Counter
from pathlib import Path
from statistics import mean

from docx import Document
from docx.shared import Pt
from openpyxl import load_workbook


def main() -> int:
    excel = Path.home() / "Downloads" / "SEA Typical School Results_2024 (1).xlsx"
    output = Path.home() / "Downloads" / "SEA Results Principal Brief.docx"
    if not excel.exists():
        print(f"EXCEL_MISSING={excel}")
        return 2

    workbook = load_workbook(excel, data_only=True)
    worksheet = workbook["Results Master Sheet"]
    headers = [cell.value for cell in worksheet[1]]

    rows: list[dict[str, object]] = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        if any(value not in (None, "") for value in row):
            rows.append(dict(zip(headers, row)))

    weighted = [
        float(row["Weighted Score Total"])
        for row in rows
        if isinstance(row.get("Weighted Score Total"), (int, float))
    ]
    remedial = sum(1 for row in rows if row.get("Remedial") in (1, True, "1"))
    resit = sum(1 for row in rows if row.get("Resit") in (1, True, "1"))
    placements = Counter(str(row.get("School Placed") or "Unplaced").strip() for row in rows)
    genders = Counter(str(row.get("Gender") or "Unspecified").strip() for row in rows)

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "Aptos"
    style.font.size = Pt(11)

    document.add_heading("SEA Results Principal Brief", level=1)
    document.add_paragraph(
        "Prepared for the principal from the SEA Typical School Results 2024 workbook."
    )

    document.add_heading("Summary", level=2)
    bullets = [
        f"Students in workbook: {len(rows)}",
        f"Average weighted score: {mean(weighted):.2f}" if weighted else "Average weighted score: unavailable",
        f"Students flagged for remedial support: {remedial}",
        f"Students flagged for resit: {resit}",
        "Gender distribution: " + ", ".join(f"{key}: {value}" for key, value in genders.items()),
    ]
    for bullet in bullets:
        document.add_paragraph(bullet, style="List Bullet")

    document.add_heading("Placement Distribution", level=2)
    table = document.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "School placed"
    header[1].text = "Students"
    for school, count in placements.most_common():
        cells = table.add_row().cells
        cells[0].text = school
        cells[1].text = str(count)

    document.add_heading("Principal Notes", level=2)
    document.add_paragraph(
        "Review remedial support flags first and prepare parent follow-up where needed. "
        "Use anonymised student references in any correspondence that leaves the school."
    )

    document.save(output)
    print(f"DOCX_CREATED={output}")
    print(f"STUDENTS={len(rows)}")
    print(f"REMEDIAL={remedial}")
    print(f"RESIT={resit}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
